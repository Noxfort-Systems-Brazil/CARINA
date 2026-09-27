# CARINA (Controlled Artificial Road-traffic Intelligence Network Architecture) is an open-source AI ecosystem for real-time, adaptive control of urban traffic light networks.
# Copyright (C) 2026 Gabriel Moraes - Noxfort Systems
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU Affero General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

# File: src/blocks/docx_subscript_formatter.py
# Author: Gabriel Moraes
# Date: September 2026

import re
from typing import Any

try:
    from docx.shared import Pt, RGBColor
except ImportError:
    pass

from src.blocks.docx_template_loader import DocxTemplateLoader


def add_formatted_text_to_paragraph(
    p: Any, text: str, font_size: float = 11.0, bold: bool = False, font_name: str = "Arial"
) -> None:
    """Splits text by ** (bold) and * (italic) to apply appropriate formatting runs inline."""
    parts = text.split("**")
    is_bold_part = False
    for part in parts:
        if part:
            sub_parts = part.split("*")
            is_italic_sub = False
            for sub in sub_parts:
                if sub:
                    add_text_run_with_subscript_support(
                        p,
                        sub,
                        font_size=font_size,
                        is_bold=(is_bold_part or bold),
                        is_italic=is_italic_sub,
                        font_name=font_name,
                    )
                is_italic_sub = not is_italic_sub
        is_bold_part = not is_bold_part


def add_text_run_with_subscript_support(
    p: Any,
    text_segment: str,
    font_size: float = 11.0,
    is_bold: bool = False,
    is_italic: bool = False,
    font_name: str = "Arial",
) -> None:
    """
    Adds runs to paragraph p, detecting variable subscripts driven by JSON subscript configuration
    and applying native Word subscript (run.font.subscript = True) to the subscript portion.
    """
    sub_config = DocxTemplateLoader.load_subscript_rules()
    if isinstance(sub_config, dict):
        pattern = sub_config.get(
            "subscript_regex_pattern",
            (
                r"([a-zA-Zα-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe\']+|"
                r"\b[a-zA-Zα-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe]{1,50})_\{?([a-zA-Z0-9α-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe]+)\}?|"
                r"\b(v)(real|limite)\b|\b(F)(ideal)\b|\b(P)(95)\b"
            ),
        )
    else:
        pattern = (
            r"([a-zA-Zα-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe\']+|"
            r"\b[a-zA-Zα-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe]{1,50})_\{?([a-zA-Z0-9α-ωΑ-Ω\u0370-\u03ff\u1f00-\u1ffe]+)\}?|"
            r"\b(v)(real|limite)\b|\b(F)(ideal)\b|\b(P)(95)\b"
        )

    last_idx = 0
    for match in re.finditer(pattern, text_segment):
        start, end = match.span()
        # Add preceding normal text
        if start > last_idx:
            normal_part = text_segment[last_idx:start]
            r = p.add_run(normal_part)
            r.font.name = font_name
            r.font.size = Pt(font_size)
            try:
                r.font.color.rgb = RGBColor(0, 0, 0)
            except Exception:
                pass
            if is_bold:
                r.bold = True
            if is_italic:
                r.italic = True

        g1, g2 = match.group(1), match.group(2)
        if g1 and g2:
            base_var = g1
            sub_text = g2
        else:
            groups = [g for g in match.groups() if g is not None]
            if len(groups) >= 2:
                base_var = groups[-2]
                sub_text = groups[-1]
            else:
                base_var = text_segment[start:end]
                sub_text = ""

        # Add base variable
        r_base = p.add_run(base_var)
        r_base.font.name = font_name
        r_base.font.size = Pt(font_size)
        try:
            r_base.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass
        if is_bold:
            r_base.bold = True
        if is_italic:
            r_base.italic = True

        # Add subscript portion
        r_sub = p.add_run(sub_text)
        r_sub.font.name = font_name
        r_sub.font.size = Pt(font_size)
        try:
            r_sub.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass
        if is_bold:
            r_sub.bold = True
        if is_italic:
            r_sub.italic = True
        r_sub.font.subscript = True

        last_idx = end

    # Add remaining text
    if last_idx < len(text_segment):
        rem_part = text_segment[last_idx:]
        r = p.add_run(rem_part)
        r.font.name = font_name
        r.font.size = Pt(font_size)
        try:
            r.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass
        if is_bold:
            r.bold = True
        if is_italic:
            r.italic = True
