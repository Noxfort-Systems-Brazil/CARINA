# CARINA (Controlled Artificial Road-traffic Intelligence Network Architecture) is an open-source AI ecosystem for real-time, adaptive control of urban traffic light networks.
# Copyright (C) 2026 Noxfort Systems
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as
# published by the Free Software Foundation, either version 3 of the
# License, or (at your option) any later version.

# File: src/blocks/docx_image_builder.py
# Author: Gabriel Moraes
# Date: September 2026

import base64
import logging
import os
import re
import tempfile
from typing import Any

try:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:
    pass


class DocxImageBuilder:
    """
    Renders Markdown images (![caption](path_or_base64)) into python-docx.
    Handles temporary file lifecycle for Base64 image data.
    Follows Single Responsibility Principle (SRP).
    """

    @staticmethod
    def render_image(doc: Any, line_str: str, font_name: str = "Arial") -> bool:
        """
        Parses Markdown image line and inserts the picture and caption into the document.
        Returns True if line was processed as an image, False otherwise.
        """
        img_match = re.match(r"^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$", line_str)
        if not img_match:
            return False

        caption_text = img_match.group(1).strip()
        img_src = img_match.group(2).strip()
        tmp_img_path = None

        try:
            if img_src.startswith("data:image/"):
                b64_data = img_src.split(",", 1)[1] if "," in img_src else img_src
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_file:
                    tmp_file.write(base64.b64decode(b64_data))
                    tmp_img_path = tmp_file.name
            elif os.path.exists(img_src):
                tmp_img_path = img_src

            if tmp_img_path and os.path.exists(tmp_img_path) and os.path.getsize(tmp_img_path) > 0:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(tmp_img_path, width=Inches(5.5))

                if caption_text:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(8)
                    r_cap = p_cap.add_run(caption_text)
                    r_cap.font.name = font_name
                    r_cap.font.size = Pt(9.0)
                    r_cap.font.italic = True
        except Exception as e:
            logging.warning(f"[DocxImageBuilder] Failed to render image line: {e}")
        finally:
            if tmp_img_path and img_src.startswith("data:image/") and os.path.exists(tmp_img_path):
                try:
                    os.remove(tmp_img_path)
                except Exception:
                    pass

        return True
